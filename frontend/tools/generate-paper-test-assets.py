from pathlib import Path
from hashlib import sha256
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader

root = Path(__file__).resolve().parent.parent
out = root / 'dev-assets/dev-files'
out.mkdir(parents=True, exist_ok=True)
title = 'DEV paper transfer test'
lines = ['Original local fixture for frontend preview and download verification.',
         'This is not a real examination paper or private teaching material.',
         'Question 1: What is 2 + 3?', 'Answer: 5.',
         'The same sample is provided as Word, PDF and image files.']
doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Inches(0.8)
section.left_margin = section.right_margin = Inches(0.8)
doc.styles['Normal'].font.name = 'Arial'
doc.styles['Normal'].font.size = Pt(11)
doc.styles['Title'].font.color.rgb = RGBColor(0, 0, 0)
doc.styles['Title'].font.size = Pt(22)
doc.add_paragraph(title, 'Title')
for line in lines:
    doc.add_paragraph(line)
doc.core_properties.author = 'DEV frontend fixture'
doc.core_properties.title = title
doc.save(out / 'paper-sample.docx')

pdf = canvas.Canvas(str(out / 'paper-sample.pdf'), pagesize=A4)
pdf.setTitle(title)
pdf.setAuthor('DEV frontend fixture')
pdf.setFont('Helvetica-Bold', 22)
pdf.drawString(50, 775, title)
pdf.setFont('Helvetica', 11)
for index, line in enumerate(lines):
    pdf.drawString(50, 725 - index * 30, line)
pdf.showPage()
pdf.save()

im = Image.new('RGB', (1000, 620), 'white')
draw = ImageDraw.Draw(im)
font = ImageFont.truetype('C:/Windows/Fonts/arial.ttf', 22)
heading = ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf', 36)
draw.text((45, 48), title, fill='black', font=heading)
for index, line in enumerate(lines):
    draw.text((45, 140 + index * 65), line, fill='black', font=font)
for ext in ['png', 'jpg', 'webp']:
    im.save(out / f'paper-sample.{ext}')

assert [p.text for p in Document(out / 'paper-sample.docx').paragraphs] == [title, *lines]
reader = PdfReader(out / 'paper-sample.pdf')
assert len(reader.pages) == 1 and 'Answer: 5.' in reader.pages[0].extract_text()
for ext in ['png', 'jpg', 'webp']:
    with Image.open(out / f'paper-sample.{ext}') as image:
        image.verify()
manifest = {p.name: {'bytes': p.stat().st_size, 'sha256': sha256(p.read_bytes()).hexdigest()} for p in out.iterdir() if p.is_file()}
(root / 'dev-assets/paper-test-assets-manifest.json').write_text(json.dumps(manifest, indent=2), encoding='utf-8')
print(json.dumps(manifest, indent=2))
